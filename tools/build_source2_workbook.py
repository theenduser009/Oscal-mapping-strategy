"""Build the lowercase four-tab Source 2 OSCAL mapping workbook.

The first tab uses the 2026-09-17 reviewed Source mapping snapshot. The other
three tabs are rebuilt from the row-by-row GitHub CSV evidence. This script is
packaging only: it performs no Snowflake reads or writes and does not change
runtime mapping approval.
"""

from __future__ import annotations

import base64
import csv
import gzip
import io
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
MAPPING = ROOT / "Mapping"
OUT = MAPPING / "archer_authoritative_sources_oscal_mappings.xlsx"

# Gzip+base64 snapshot of the reviewed sources_source tab (A:S), generated
# from the accepted 2026-09-17 review artifact. Keeping it embedded makes the
# binary workbook reproducible from text committed to GitHub.
SOURCE_REVIEW_GZ_B64 = """H4sIAAQHrGoC/+1d63LbyHL+v1X7DlP+s3YFJGVJvq2yqaVJyEZMkTy8eNfl2mKBwIjCMQjw4CJZJ5W/eYA8Yp4k3T0XDEBQJnWxaSepk1qLBDEz3d/0vWf+iJNP6QXnGUviK6udeBc8YecBD30WuUtuDcaddo8tY5+H8t885EseZWzlZhfWmbtaBdGCZdcreDYJFkHkhqwfZzy1Om7i459Bdq1e7LuZK56dJG6UekmwyoI4YmnmZjn8InST4DzwXPow4tznvmVfBj6PPPmT8zhZiq/DeBF4ck7Fe9WEEn4Z8Cv14lEeZcGSs0y9ggV+8aGbLLhcT5d7QYpvb8HwnzPmejiWNY7zxOMsvXBXnHE1oZEYI+HnPMEPfv7p0BoPpqOOPeu3z2wgQObCLFmDnfHMxTlanviouZQfNLMgC7nVDRLuZdYwCZZucs3oQxafs+yCMy+OsiQOmfypNYblwU/et0edt+2R1RmMbDax/5ywTns4mY7srmWND56yf2Hjg0OgWbBsXQIPYDDOojiah270iWWwthO2dFfMp5FDGDNmalJi/KYFc/mUrxphEMFa7Xb3g4W/27QGoDzNVw8ieZ4K2iGaSqPhs4pC1ZHHWQKrPGGreJWHMHOfBcDw4BJ+MR21J3Z39u/jQZ+l7nKFj19k2Sr9tdVauQueNqMgzZqL+LJF2GiMFHtahOK0dfm0edg8aslltP6exlFD87D1809Hionv7dHYGfS34uMlTxA3ipPvxZ8MkZIBpHFTxYng59fjo5xU00qJnjcyUS1AsVGyTX68kXN+7OUkDvRY98i7VsjdJGpleRYngQusW/CIJ27YUrMGbh0rbnXtcWfkDCfbcizhiO9UcazLC3Ekd56bZxcwMIgQnL0gyAnwL2Iwl5jIDtRQL5673qfGx07PGQ7t7l9V9vZtuztmYqrEZmucr1ZAz/M8DBWxkTO/srIYbRooGCY85YmeCz3P8si7cKMF0BnlWcjGZzbzLuI45SmTa2SXKcNZe1kOr4WXiJ/7xZL1Zl/mIbwEdjx7DAv1CW9PrK59ao8AkkBQeL1lWY9e5yk8lKYANjdCgRukzEN2WWyeZ6BBcAEwQAb/IJrC9/AHrDVe8YjNuefmKWdO/3QwOmsj0/AJl6V85SaEcTdJBOFJGzVZNwbcZywGoF0BU0A+Jigbgfw8wh8At2JUDZ4bhtfNRxKHxnsAhinSPABSVdG3CLKLfN704mULOM8jHyaXHBy8ag1SeF9jKbRKA/YRDLW4bs3DeN5KA3hJcH7dUJhqzOM88kGCt/6R49LjKG0BvWaDod2f/W1qj3Gd4+bS//mnZwq2k1G7887pv5k53a1gu0ri1cffUZr+9gvMxvuE8wr8X/4q9KM1jQKYgCF8UEZck0JFkqk9DeCdTp0uEzoV6ZnCRgNKpoxGqWK5a3ecs3bvfrDcl4N7MU88RCCIFIm3X5mQV0BagK3P+DLIcFK4aJ+mxpMMvjrPQKyqPxtCwUj93bT0YyA1c/3KG1DdGk37E4fQPSR1AK+m7Ub2zYVJTpwM0p678wDNG2F/ADRQ32WwGyWVM9oYGRLaFau9ApzFsD+UDdFk9mfYlsWiaBUkXpBbKSOJF3gE+wR+QxCOk6xp9UHu4jcIUDfZL3Q/V+juOf13YwPXrwGw7MzNgHNWSXAu6bOmkk0ff8/zwP+rmYAo+tS8AH4oOW1/hucQUfhNiqRV0kxqIldYbLtoVydKVxy5oFUWGMMpiPqEk2TRIrOFg7I4Qqwn8RKFKv8MVPKAydNRr6VFKxltwNimlSdBYV2AWakfeQxyEOzjRRmHYlON37aHCMROHJ0HYK0WEyM0g6CEbQoDpq3pyCkM0BTk6jnuYJo4gs+gtx45xZ1OhE0RRbBl17Fjqm9h9aLsvQYzOZ6TCvIfDl3AyLQFpqP3aRUHUZa2Dg8OnzcOXjWevmiIFRzK/za8HN/pN8iGk6hH/L2wDMVyG2ug7fsB4shF4hQ+hzvHzVuYR7vBTKvvIPI5ch/QGl6fAHKJul4chrjDUUawTROsCjbkpVDdGwAl1LWxHgk/c1nLPM3IVkAgKAVMAF+3rQw5LTR2WSoL9Z9dwFfwP1wWsD8B6eUjekAtbwDcNxRWL62uM+702iD7RzurYD9IvdAFNzIBDYyyKSLVOlTk6PEF0ByYlKdg4LLi8XQ37DxCD0AYfFolaiT8x9pcLIGO3/4Vf/Fv/wl0r9eHN7oE+JP0IyxL6l933V8r6S3fWB7AJE8SdAyUBpNbVsQW0jBfgJoDMeVfs2v8KDV/XO9EbITM/fh9r6zJYOh0ZiPcPXa/Y99FcXmBVEQ1mBgV8hrUV8LFArN4FXgkm++gyUYwjRAkjF6Xzyb4Xmm8ZAHqDsBPfAVMSS+CldIYiQyGyO3VtAqdRRPEDYSPF0JmZJ+1h2xk/23q0CQAoiAXw8JVIvEQqTCAl8RpamhCNMtBWpDgk7QieDfZBJWXECEAOhF6+Z//+m+xjtJklBWlRShtCwFWU/OpAWAbtHFYdPDi+d9B0KekqUHuobIEbQu8yhyfSNTjlzx09lvPPT2w2pNJu/P2zO5P7gLXuZvy58dK+70Gkz0BAzvLXO8CkUickM/shEfptGXEUjeZB0Cc0psLUQIjiGCeGAcRKI0u8YHk2FZ20yNlOF1dcDLKaw2oOS1Tvd4T3LfqJmcxp5tauDfR6NrOytpLXff0KTCsPxkNerPxpN3vtkddDRsNlEUS56DgSIyJ4CP9UcgtMJj9APYeRhNUeDLNXJwDyHEyHGj7SiPpDGMK7HEYfOJgOGsR8OS2Mk0Js2slypAHNGnFGZQDxBk5OzCb5L9axVtqgfTWsUc4qw+Aoj76y67pFejfCkbS+GgNCiHVIhHVGnMKGrfG+bwh/23IW5+nwSJSE12FrmfOE9GidjGtaLOJ/i1RdLiOotnpaHA2Exp0bHfw4Rl8NRtPX6s/e/Z7u3cj3L6MvD5MTUQxJOqS3Mty2HTCbL4A79wFhYNRIEFANHLuCEApxPwcPT20ezFwBmR2o6xgnAi+qfGvDY7DRvD5eRChz3QLFHaLweCXYQNQnoKpDsN5+GL8OUPdCTA8Bxg2MnceFvFZN4WH57Bqgbu9RNNRDZq+Epy6mqfCLm2gu8LdhCyzOe1ziodtgHsRIvsixO4raKbDnMBRhUiSMAmAFBy6MPinsHUsCSwTV+2e0x4DpsAoTdwGMAgg6qEcEqs3/UCwxkGCepkEtpESIGQj8FHDplrFouUlwsxh4KYtvVv2EnLH1ll7PLFHM8nY8aw9nbwdjJxJe+K8t2eCT+vgMmHUxBDKx9/BIv3tF2nE/2Jia2SaqpiWcdNMEI5EF8kEcsrTe9eQfqEigUngFlk41y9JnYoRtab+qia8qQRFDlWPWkioqq2k3J2b40/fEhvPtDgajgYduwvULkxr2GadeLmKIzQPu5qF1qNahKzcJFORAkxFU+YcQAIgkO9oFDBo6g/VSxq4DK5dwab+E1yBhP8jB2sd/2rqVz8qMskxoID0IuYd9Q9NW+NWlrzILIHC4aIogKVBSGEslUH6IsgoIPWHLjtA4gBKYlDeWmcBoOqI3GSnJBbFwIWCJbE1F8F/EIraCPzMvVyIwj1E2XNwoHuUDt2MtluJnU55l5ZiDCsNiwc0ym8vcdbD3Wsih2LR6xIn4N+jpHlh9adnr0ELDU5n6xaGSqGgmbNzXFJJEA9mk9WHJjv4FYZCtA0LPrrU8SEGP6p5v5sR0REpL0a7lau0FIBPePU6UigTceSHn5QDiNVpN2+MWgItZMqO2f237X7HxiiIJSoVNkczscBBTo5GUenOIv1WmkSTjYSkLXJxOuVTzj3qFJ1Kbq9HPNEVCM4JzGsZvBMZ0D/YM5C+BL6fDcFwBALPRmAe9d9YbfAq0pRCJCAIQJKklqs/gu1KHzXFf5uiFAtUlnbGFYFTHgr/eO2Zj79rFYhqtARsXIsLwqCR0D6vB/fgEgtFQlY8zsTjug7nNhkcJfVCdw4a6DJwVXQzjGNMwVBmUOSl10lUaKzWIkETO0iVFY2uYY25fjbo2j3WtcfOmz5GtUplaJR8Icg0QrDVVeKFEypBR/t1U6jSGSYu/YdEoXyNOSIy9jkTEdYU/csAdLfPYOGcQqpEgmIqKlJHHvl7irf14GPHT/c7mPoKuA8CxZDG491LMmL4UiO9HpsTfIZF+XIOZDPlbxDVVYh9RfFbmf1Dy18abl3+lmexuwD+TiTr4UGBt/6gr6TspADfTmLWRGEUR0rwZTeZAH2NQfhFQ/9kzUXYDoLV8jSJMZ+KaVSCSQmcbQRee7FI+EJWPBqSHIiN3AUZHEcLMnDrpG0hi3FTCSq1dL2iFsE6Nvd9gOapNew8OGQAHR56nQtej5uh/v47wE6xmPtGzLCmRsZUkPuFnEPrG0qadimsqoODrQx46xEJjaaDnZDzjvMVaASsovvEr2W4sj5UKWtVbwyNTtSEighmA607VJryrTCMzjgHEQLUiIHiGKn0LMDSLMC3jx7o4ZGOQkwGs3f2hzXDpz4EYSIAqKF0dT3rMVkIb4lJz8DTLS8JqEz3lqKiy2GHLrFEWZFdlRu5yuUi95W5SgS0zkOXQncu8kPbvmYtQYGJ+roGgU6d05QjSDsGqX5Str1VeAInpa0brCBjWGqITS5q9cJwj2IdGNaBFPQPXLBoMKR1LTOqNCCllr+9sroHm/vw2Dp1+l3wKXeUQSCBkVuyfsEIfxlvkc+kRonNg+aldQhKF9HIGdwm8L45BqaqXigMdpObuRYRUwTZS1H0TCNhQy7mNvhoStY3RNmqqKgqwHKqEKI3ZRbXtn0wI9x+59iptnWErlIgMaZYLnPaNU5QTTwptpsjpLK+XPaEyG1qBFpVAmcvsfJcp2gm+Hn/zWxkj6e9yXh22nZ6oM2A9Pao3+7psKqU5+Sh3gpIWJ0Lfi1woOTuw0zlL+qV36kbhEa9QkYdMQtpW6b3VJkgjBA5U7BA0KkrK7XNYCFhk5qBpQ1iZXNKTw0s9VzA9xM0L4AK7x0qCb9PEaNK6EpVwSCjaZvVg6IrXQqsmFVPyropXUZ1b0JGmMZkv1T9Ka2nvgyTrpon9VzxlShDDgAXtVFGap2RaUtVUQB8aSwwRiO97b0DyMsvAoT6WYSxvF5Gd1/6CT/jl+hP6S4hTGEZaqtHrS8GdLALpmJWPhiSquaIUGJVhQVsxlYpf2sRpFfTyOJaC72we1TJcpVMojtHg3oPEfZKd/19GNo7x5alksY+93q58mgCX6GXVWvCPOZNoAqQbZELqloaIxYIHxjiKk4+PXl0u5RIHGGhfYKWA+pJkSJJi2YFGXcuB5TTL0aYe4PBu+nQcinHIHMSm2PLajIyJVHKPYDpn8T54oLhsuzR7MyetGfv272pXdNRQfFkg97NHybLcXRgdUCsOZ12z5l82D2vLH138E83ZJWLB0QWWXVT75JvOyJwHWt+7i2mNiDHoNIPhJynSnoZABI76C4wahBhNuYnKMphPM6oYHG7iJGC0cbcWC2A1osS1ib70IkxUZS5XphQnYiRG/sxcmJHhxYaXtreOnV6k1t05Z0HYYaBPpxVnGySVOJbWS+L4UHxM6qRwxMGksC9V3FV6V+vTHJLsaXb1B9QfFVnJpvU6TtdJ7AJZz+MuDuyhoOe03Hsemeg1CdYVxL6hRpPo4xvFcPWvS5V8fVUd7su3aNndG/gra36KlbvtYCvErksfqy67RBwciV6uGo4oY6Yqojvu8q4HR0LBH0QNXyzo1vhxJRrgnQi+VAv1N6aDTCS0qVivu3Qoap91QEIejpgxoVcHX+Q6cQIibptfL1a5hYRa8SHbn9NrhvFuEXPV7LO9qLlbKzCtaBB9zLKcPSsAoodOl0eHD1DAzCCqKWGKmpRLnqudoQUhZ6qbVJ1AgE5fHPAe0N/lHyRapMSzVFpTXdUefwCZBqIe4mc51Xk7BN0tu2kcudgKuyGnFLX0x2anGQc5Iu9Thp7NN/972Y6emEVcUkJkXLx+N3ME2kobbBSiiN4ZFKcngv4nS2UIs8/rx42pkIKhp2tzQhsUMFn3EiWPhqnkN2U5e+tH1F3EV+ZRxgYJq08NYAn4gQDtxoVLY6/2GjNqEClOidFHuykkv28lJT5gcpsj14aXQ/awr4zXmuk1Ta9Dwqst+t9UFDVIXJRmXgiW5Fk/ZHgkziGRBnXW1d7m5vqWtYV4WFuIFKBl+d5WFP4XYu5Qsupem4JPMOQ+nEKS45eWXb/TfsNhV5m4w6IzZnTL7I2BDczZTMM3cjM16zg7yYWHIYG3tImHgh7iWG9UmDeizeF5Mf4lSjy1uWrPFq4CyoluF14Qcg72vYb2w/UPMt96DuVEdzUbXDCKqQrBlRmXNHkoI4Ok3WT1IEhX37448QLjg/MRpkCfGOJvqK9T5htY7sHNhromi+moneFZprx1YZG0FoUprvBsP60vBsE281BAn1AZ4uYL+C2CvN0M8S0B1IpuBIhXb6fh2YcP7XUB7Oe8xr8lw/3JX/U2LB95ni+zoYS3GIw9YPIDRKzgO2uJtsGaKCTtd6bTnUiO8NFnd9TLlIhfJRXtVavsseJ4eNDy37vdG1qsbOHg7EzGZThsVNFQU31QOkMTWW264x6wsFwCLI4ub5PUZAnwd3ih8L2KKqS1Hwb+hDNjUHEmqqU7zGEeHxk2aen6MuDauiCa7JzKoSfn2Mb3yVv4JGNpWOCu+itS0LOOSwK00jq8RN1hi0CBZ+E/3fpiH2JEExCjCfts+GWyTc6vZw54wHD2iMg+HLVolMk4f0f4P8aZ2eNbreSi9t8ji/9NOSfMTD1ROZInFM2HAynVNBu4QOb0yKPtMluuFp68QxeijElKu2uScSVidpkE1mALt9DB15GqGWFeyj9UNVIssrn+tgWpMXmk8JOGJZkzMG6Q48C7HHgkKb7PZydfmwNTk/BAWtjgeRE1kfeDmcJz2R97DrQBufn4GThgeAG4gKsE8Pf+PcOtRKi1tK51ZnuhrENud0bwVZtbipmsAFflSnuloD7Rlh6Zo0mf94dRtnnxo1QShZuJKN/DV3zViHowwLqa4kojZp4mzXXgGidlN8MHM8tRMOsA0TAhW+FilxqZIP7DprB2DhEah6vpajj9gRk6h+gqfG//4yj23NfH3RcCHRNY90Bhic9bJ585VoCYnmDim+Nd950qGrfGU9YzXvF5Q1zlAOGOpGB0YrCCSIs9FPayFB4MDGscgibTINNH9fKCwWlD4UOsb8uxEgs3jKhSk7odDl07dNiyKJ/4JE1Ufq+gX6F2h7qKPxMcgkgCWri884H/W6FvxdWrz2ezKbD7tb4C900ayxjn8pYDQz24HMmPve+PQjFYtGaFTwvzRoj2HmIAZcwjK80cxryR3TzUtHnoaHZQNY0NF+kwhFxm0iUlOqDUdcAW8ZreToFZm8AKd3ZxPV9H6X9A8iYg8EMvBZ2W/mKIZH0ZflKmJVoWe0D9l5asuIUg4+oGOtLy2UE/NQZAVCH09c9Z/z2DqLyFH2wkmwIVCLjwbD66D1WfV/X9GWO2n/g+fb5MrKA+D6XeVhKrlmAiyS4VC3AdMyuO8f2ATy2fGN8oDfoiOP9qUMTRsHWbryzAxip24mvXBF+9PkSvKxM7BRxlLmZ6cHGYXn4RoX/FjmeEvWF9C2CC+JSG7paBb3SMJcHTffj4q6EzSPukaf5aieY3q84lVv24REqAaoyPerwG+TxWrxomy6ZdSzeAxRPFN4qwlxPkO6HkbnrPEpEWgCrXb4zzD072Alz8BBIxzNYjtOfiR/e5lw0kQFZdy3k6XOEiJboX/kaCr4ektROHqkriMi6L66W0dMsLP7gXF+0sckhaRW2502xOFOunm5GLh5UbYCryZxz3fRjycNg18IhotynetiawRBwQ9BgIXGqK5rhOTp9Qd589P3J1mdPwU4CFHSxW2wyHd8iuIKWTUPcZrmh92dMXyoPQNpCjylfgFoMa7gukTU885rbNvmsFTaL8nSne0MzzRXi1r10gxBV/EkBYZG2NN3UG/3nO7ZrmOahoIwX+1zPG78Gz1x044p3ipJ6lUvdUBpd4sTaJVw0xOPDJ99btvzZIQqC9rQ3mUmgDu3RmTPG6y+3Q2sSA5jqkk6eh+6iauij/Y9VPWl6l4OPFS7FVSFaNrZSDivEdIC+VUK6q8bFgsLtwPm28Hje4qxifWkXii6zKMISJT74ExzLE5U9MF+Vj/A3lRMZF6EUSUuTAvpSizd4VDmmuVtTQAP+44kZO6axxfUo0i/C2yfF1XK1vBAXIKoBzYy6HsmcBh6QkracrkisqSnsN2bFsfFU3iGqOW5zRiq6w2KgDYVCeQr+p6FnxWE2VIRGx7Sz9DrN+PLh+9oep0/2obOtRLMfp3rj2bHl9Mf2SIat+602WkTE2gJNVvUDe9JrkZi5chN+EWMoUFW/UqTteiXjFrgkcQAN7eddbcQRx2pDUd+Dfir3LiJ8c8vN/aC4pqaJ87bsPzu9addmWGUih7NopsUkDYtM3q5Kp2jFeZYGPjeiMyv3Ooxd6VpUPRddlilBUbUJvwvN98wSXuz/8/3/Ft+fW2C4dt6Np2eyhXVL1qObBwuGeSaiV2oHtm+rHO7K9D8Uwwmkjcpsv8T0GpvJwI28vwhNGu1v/rAgebFrfIIeusVNzlJjK+1K9znXnWqtboZRWRpF2/I188Yte3hSvnFT8W5lzY/Ap68cH0xZe5EPAKfSxwtC1JkuTPdFkgcBTqKs35FRACzpW+tSJqvihMUYL74KUjrVgw7KK4RW9bUbL7Nkj+OVuOS0FNdQ+Xpd2DpZu9vdPHJP3UqpghkI5iJ8zQomi6gh3phNz6ytHgDfdc7AYBe14Coori/tEXdDBxHtQkFmjGVUkWCSvGnhTIwpwHfEPZxKLm79dqkVgT09filvMS4HI3WEJIq/xmW+X4yN/C9CeE/agYQAAA=="""

SHEETS = (
    ("sources_source", None),
    ("sources_topic", MAPPING / "SOURCE2_TOPIC_MAPPING.csv"),
    ("sources_section", MAPPING / "SOURCE2_SECTION_MAPPING.csv"),
    ("sources_sub_section", MAPPING / "SOURCE2_SUB_SECTION_MAPPING.csv"),
)

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
READY_FILL = PatternFill("solid", fgColor="E2F0D9")
DEFER_FILL = PatternFill("solid", fgColor="FFF2CC")
REMAP_FILL = PatternFill("solid", fgColor="FCE4D6")
EXCLUDE_FILL = PatternFill("solid", fgColor="E7E6E6")
SUMMARY_FILL = PatternFill("solid", fgColor="5B9BD5")

SOURCE_WIDTHS = [10, 34, 24, 46, 20, 44, 20, 18, 22, 34, 18, 46, 28, 27, 20, 34, 54, 42, 48]
DEFAULT_WIDTHS = [10, 34, 24, 46, 20, 44, 20, 18, 22, 34, 18, 36, 24]


def csv_rows(path: Path):
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        return list(csv.reader(handle))


def source_review_rows():
    raw = gzip.decompress(base64.b64decode(SOURCE_REVIEW_GZ_B64)).decode("utf-8")
    return list(csv.reader(io.StringIO(raw)))


def style_sheet(ws, widths):
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(index)].width = width


def main():
    workbook = Workbook()
    workbook.remove(workbook.active)

    for name, path in SHEETS:
        rows = source_review_rows() if path is None else csv_rows(path)
        ws = workbook.create_sheet(name)
        for row in rows:
            ws.append(row)
        style_sheet(ws, SOURCE_WIDTHS if name == "sources_source" else DEFAULT_WIDTHS)

        if name == "sources_source":
            for row in range(2, ws.max_row + 1):
                status = str(ws.cell(row, 14).value or "")
                cell = ws.cell(row, 14)
                if status.startswith("READY"):
                    cell.fill = READY_FILL
                elif status.startswith("DEFERRED"):
                    cell.fill = DEFER_FILL
                elif status == "REMAP REQUIRED":
                    cell.fill = REMAP_FILL
                elif status == "EXCLUDE FROM OSCAL":
                    cell.fill = EXCLUDE_FILL

            ws["U1"] = "Source tab review"
            ws["V1"] = "Count"
            ws["U2"] = "Total mapped fields"
            ws["V2"] = ws.max_row - 1
            ws["U3"] = "Ready / ready with condition"
            ws["V3"] = '=COUNTIF($N$2:$N$57,"READY*")'
            ws["U4"] = "Deferred"
            ws["V4"] = '=COUNTIF($N$2:$N$57,"DEFERRED*")'
            ws["U5"] = "Remap required"
            ws["V5"] = '=COUNTIF($N$2:$N$57,"REMAP REQUIRED")'
            ws["U6"] = "Excluded from OSCAL"
            ws["V6"] = '=COUNTIF($N$2:$N$57,"EXCLUDE FROM OSCAL")'
            ws["U7"] = "Not yet reviewed"
            ws["V7"] = '=COUNTBLANK($N$2:$N$57)'
            for cell in ws[1][20:22]:
                cell.fill = SUMMARY_FILL
                cell.font = HEADER_FONT
            ws.column_dimensions["U"].width = 28
            ws.column_dimensions["V"].width = 14

    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)
    print(f"Wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
